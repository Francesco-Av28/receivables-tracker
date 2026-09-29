# Crea risorse\modello_vuoto.xlsx: stessa struttura dell'Excel dell'app (fogli, intestazioni,
# formattazione, formule) ma SENZA clienti né movimenti. Si lancia una volta sola, partendo
# da un Excel dell'app esistente passato come argomento.
import sys
from pathlib import Path

from openpyxl import load_workbook

QUI = Path(__file__).parent
ORIGINE = Path(sys.argv[1])
DEST = QUI / "risorse" / "modello_vuoto.xlsx"

wb = load_workbook(ORIGINE)
for nome in list(wb.sheetnames):
    if nome not in ("Riepilogo", "Clienti", "Insoluti", "Acconti"):
        del wb[nome]

# Clienti: via tutti i dati e le formule (l'app le riscrive quando crea un cliente)
ws = wb["Clienti"]
for r in range(5, ws.max_row + 1):
    for c in range(1, 14):
        ws.cell(row=r, column=c).value = None

# Insoluti: restano solo le formule automatiche (B, F, G)
ws = wb["Insoluti"]
for r in range(5, ws.max_row + 1):
    for c in (1, 3, 4, 5, 8):
        ws.cell(row=r, column=c).value = None

# Acconti: resta solo la formula della ragione sociale (C)
ws = wb["Acconti"]
for r in range(5, ws.max_row + 1):
    for c in (1, 2, 4, 5, 6, 7):
        ws.cell(row=r, column=c).value = None

# Riepilogo: intervalli del foglio Clienti estesi a 1000 righe
ws = wb["Riepilogo"]
for riga in ws.iter_rows():
    for cella in riga:
        if isinstance(cella.value, str) and cella.value.startswith("="):
            cella.value = cella.value.replace("193", "1000")

DEST.parent.mkdir(exist_ok=True)
wb.save(DEST)
print(f"Modello vuoto creato: {DEST}")
