# Collaudo end-to-end dell'app con il client di test Flask (fatture finte).
# Azzera da solo i dati locali di prova (_prepara_dati_test.py, vedi percorso_dati.txt): si può rilanciare quante volte si vuole.
import runpy
import sys
from datetime import datetime
from pathlib import Path

QUI = Path(__file__).parent
sys.path.insert(0, str(QUI))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
runpy.run_path(str(QUI / "_prepara_dati_test.py"))

import pdfplumber
from reportlab.pdfgen import canvas

import motore
from app import app

CLIENTE = "20"          # ALFA ESEMPIO S.R.L (P.IVA 07654321004), cliente finto di _prepara_dati_test.py
SENZA_FATTURE = "10"     # BETA ESEMPIO S.R.L.S (finto): nessuna fattura registrata
esiti = []


def fattura_finta(nome, numero, data, importo, intestazione=None):
    """Crea nella INBOX un PDF fattura finto (per ALFA ESEMPIO S.R.L se non si passa un'intestazione)."""
    pdf = motore.INBOX / nome
    c = canvas.Canvas(str(pdf))
    y = 780
    for riga in [
        f"nr. {numero} del {data}",
        "FORNITORE CLIENTE",
    ] + (intestazione or [
        "STUDIO DI PROVA ESEMPIO ALFA ESEMPIO S.R.L",
        "P.IVA: IT01234567897 P.IVA: IT07654321004",
    ]) + [
        "METODO DI PAGAMENTO",
        "NR RATA MODALITA PAGAMENTO DATA SCADENZA IMPORTO",
        f"1 Bonifico 19/09/2026 {importo}",
        f"Totale documento {importo}",
    ]:
        c.drawString(60, y, riga)
        y -= 24
    c.save()
    return pdf


def testo_pdf(percorso):
    with pdfplumber.open(percorso) as f:
        return f.pages[0].extract_text() or ""


def ok(nome, condizione, dettaglio=""):
    if not isinstance(condizione, bool):
        print(f"   (esito non booleano in '{nome}': {type(condizione).__name__} {str(condizione)[:60]})")
    esiti.append(bool(condizione))
    print(f"[{'OK' if condizione else 'FALLITO'}] {nome}" + (f" — {dettaglio}" if dettaglio and not condizione else ""))
    return condizione


def controlla(nome, risposta, atteso=200, contiene=None):
    corpo = risposta.get_data(as_text=True)
    ok(nome + f" (HTTP {risposta.status_code})",
       risposta.status_code == atteso and (contiene is None or contiene in corpo),
       f"atteso nel corpo: {contiene}")
    return corpo


def fatture_cliente():
    return sorted([f for f in motore.leggi_fatture(motore.apri_excel()) if f["codice"] == int(CLIENTE)],
                  key=lambda f: f["riga"])


cl = app.test_client()

# --- login ---
controlla("nessuna password: il Riepilogo si apre subito", cl.get("/"), contiene="Riepilogo")
controlla("vecchio indirizzo /login porta al Riepilogo", cl.get("/login"), atteso=302)
controlla("dashboard", cl.get("/"), contiene="Crediti aperti")

# --- smistamento ---
fattura_finta("fattura_prova.pdf", "FPR 99/26", "20/08/2026", "915,00")
corpo = controlla("analizza inbox", cl.post("/smistamento/analizza"), contiene="FPR 99/26")
ok("anteprima mostra la fattura pronta", "ALFA ESEMPIO S.R.L" in corpo and "pronta" in corpo)
controlla("conferma", cl.post("/smistamento/conferma", follow_redirects=True), contiene="915,00")
controlla("dashboard aggiornata", cl.get("/"), contiene="FPR 99/26")
pdf_dest = motore.CARTELLA_CLIENTI / "20 - ALFA ESEMPIO S.R.L" / "FT FPR 99-26 del 20-08-2026.pdf"
ok("PDF smistato nella cartella cliente", pdf_dest.exists(), str(pdf_dest))

# --- doppione: la stessa fattura rimessa nella INBOX viene rifiutata ---
doppione = fattura_finta("fattura_prova_bis.pdf", "FPR 99/26", "20/08/2026", "915,00")
controlla("doppione rifiutato in anteprima", cl.post("/smistamento/analizza"), contiene="Già registrata")
controlla("doppione non registrato", cl.post("/smistamento/conferma", follow_redirects=True),
          contiene="Niente da fare")
ok("doppione rimasto nella INBOX", doppione.exists())
doppione.unlink()

# --- pagamenti rifiutati ---
controlla("pagamento a cliente senza fatture aperte rifiutato", cl.post("/pagamenti/registra", data={
    "codice": SENZA_FATTURE, "importo": "500", "data": "", "fattura": "", "modalita": ""},
    follow_redirects=True), contiene="non ha fatture aperte")
ok("nessun pagamento registrato per il cliente senza fatture",
   not any(a["codice"] == int(SENZA_FATTURE) for a in motore.leggi_acconti(motore.apri_excel())))
controlla("pagamento oltre il dovuto rifiutato", cl.post("/pagamenti/registra", data={
    "codice": CLIENTE, "importo": "2.000,00", "data": "", "fattura": "", "modalita": ""},
    follow_redirects=True), contiene="supera quanto ancora dovuto")

# --- acconto: timbro + copia originale ---
controlla("pagamento 400 (acconto)", cl.post("/pagamenti/registra", data={
    "codice": CLIENTE, "importo": "400", "data": "", "fattura": "", "modalita": "Bonifico"},
    follow_redirects=True), contiene="acconto")
ok("timbro ACCONTO sul PDF", "ACCONTO" in testo_pdf(pdf_dest))
originale = motore.ORIGINALI / "20 - ALFA ESEMPIO S.R.L" / pdf_dest.name
ok("copia originale pulita in ORIGINALI", originale.exists() and "ACCONTO" not in testo_pdf(originale))

# --- scheda clienti: incassato e da incassare aggiornati ---
corpo = controlla("pagina clienti", cl.get("/clienti"), contiene="Da incassare")
ok("clienti: incassato 400 e da incassare 515", "400,00" in corpo and "515,00" in corpo)

# --- modifica fattura (totale 1.000) ---
fatt = fatture_cliente()[0]
controlla("modifica fattura", cl.post(f"/fatture/modifica/{fatt['riga']}", data={
    "numero": "FPR 99/26", "data": "20/08/2026", "scadenza": "19/09/2026",
    "totale": "1.000,00"}, follow_redirects=True), contiene="aggiornata")
fatt = fatture_cliente()[0]
ok("residuo dopo modifica = 600", fatt["residuo"] == 600.0, str(fatt["residuo"]))

# --- annullo acconto: residuo e PDF pulito ---
acc = motore.leggi_acconti(motore.apri_excel())[0]
controlla("annulla pagamento", cl.post(f"/pagamenti/elimina/{acc['riga']}", follow_redirects=True),
          contiene="timbri tolti")
fatt = fatture_cliente()[0]
ok("residuo ripristinato = 1000", fatt["residuo"] == 1000.0, str(fatt["residuo"]))
ok("timbro ACCONTO tolto dal PDF", "ACCONTO" not in testo_pdf(pdf_dest))

# --- saldo completo: PDF in SALDATE, poi annullo ---
controlla("saldo completo", cl.post("/pagamenti/registra", data={
    "codice": CLIENTE, "importo": "1.000,00", "data": "", "fattura": "", "modalita": "Contanti"},
    follow_redirects=True), contiene="SALDATA")
pdf_saldata = motore.SALDATE / f"ALFA ESEMPIO S.R.L - {pdf_dest.name}"
ok("PDF spostato in SALDATE", pdf_saldata.exists() and not pdf_dest.exists())
ok("timbro PAGATO sul PDF saldato", "PAGATO" in testo_pdf(pdf_saldata))
acc = motore.leggi_acconti(motore.apri_excel())[0]
controlla("annulla saldo", cl.post(f"/pagamenti/elimina/{acc['riga']}", follow_redirects=True),
          contiene="annullato")
ok("PDF tornato in CLIENTI", pdf_dest.exists() and not pdf_saldata.exists())
ok("timbro PAGATO tolto dal PDF", "PAGATO" not in testo_pdf(pdf_dest))
fatt = fatture_cliente()[0]
ok("residuo dopo annullo saldo = 1000", fatt["residuo"] == 1000.0, str(fatt["residuo"]))

# --- fatture omonime (stesso numero, data diversa): pagamento unico e annullo ---
fattura_finta("fattura_prova2.pdf", "FPR 99/26", "21/08/2026", "915,00")
controlla("seconda fattura stesso numero registrata", cl.post("/smistamento/conferma", follow_redirects=True),
          contiene="915,00")
ok("due fatture FPR 99/26 per il cliente", len(fatture_cliente()) == 2)
controlla("pagamento unico 1.915 salda entrambe", cl.post("/pagamenti/registra", data={
    "codice": CLIENTE, "importo": "1.915,00", "data": "", "fattura": "", "modalita": "Bonifico"},
    follow_redirects=True), contiene="SALDATA")
ok("entrambe saldate", all(f["residuo"] == 0 for f in fatture_cliente()))
acc = motore.leggi_acconti(motore.apri_excel())[0]
controlla("annulla pagamento doppio", cl.post(f"/pagamenti/elimina/{acc['riga']}", follow_redirects=True),
          contiene="annullato")
residui = [f["residuo"] for f in fatture_cliente()]
ok("residui ripristinati sulle fatture giuste (1000 e 915)", residui == [1000.0, 915.0], str(residui))
ok("PDF di entrambe tornati e puliti",
   all((pdf_dest.parent / n).exists() and "PAGATO" not in testo_pdf(pdf_dest.parent / n)
       for n in ("FT FPR 99-26 del 20-08-2026.pdf", "FT FPR 99-26 del 21-08-2026.pdf")))

# --- nuovo cliente dalla INBOX: P.IVA sconosciuta, nome su due righe, CF del titolare ---
fattura_finta("fattura_nuovo.pdf", "FPR 200/26", "01/09/2026", "100,00", [
    "STUDIO DI PROVA ESEMPIO LA BOTTEGA DELLE PROVE DI",
    "ROSSI MARIO",
    "P.IVA: IT01234567897",
    "P.IVA: IT99999999901",
    "C.F.: PRVSTD70A01H501Z",
    "C.F.: RSSMRA80A01H501U",
])
corpo = controlla("anteprima nuovo cliente", cl.post("/smistamento/analizza"), contiene="NUOVO:")
ok("nome del nuovo cliente letto su due righe", "LA BOTTEGA DELLE PROVE DI ROSSI MARIO" in corpo)
controlla("registrazione con creazione cliente", cl.post("/smistamento/conferma", follow_redirects=True),
          contiene="Nuovo cliente creato")
nuovo = next((c for c in motore.leggi_anagrafica(motore.apri_excel()) if c["piva"] == "99999999901"), None)
ok("nuovo cliente in anagrafica con P.IVA e CF", nuovo is not None and nuovo["cf"] == "RSSMRA80A01H501U"
   and nuovo["ragione_sociale"] == "LA BOTTEGA DELLE PROVE DI ROSSI MARIO", str(nuovo))
ok("cartella del nuovo cliente creata con il PDF",
   nuovo is not None and (motore.cartella_di(nuovo) / "FT FPR 200-26 del 01-09-2026.pdf").exists())
ok("nuovo cliente ha da incassare 100",
   nuovo is not None and motore.situazione_clienti(motore.apri_excel())[nuovo["codice"]]["da_incassare"] == 100.0)

# --- nuovo cliente persona fisica: solo codice fiscale, nessuna P.IVA ---
fattura_finta("fattura_pf.pdf", "FPR 201/26", "02/09/2026", "50,00", [
    "STUDIO DI PROVA ESEMPIO VERDI LUIGI",
    "P.IVA: IT01234567897 VIA ROMA 1",
    "C.F.: PRVSTD70A01H501Z C.F.: VRDLGU75B02F839X",
])
controlla("persona fisica solo CF registrata", cl.post("/smistamento/conferma", follow_redirects=True),
          contiene="Nuovo cliente creato")
pf = next((c for c in motore.leggi_anagrafica(motore.apri_excel()) if c["cf"] == "VRDLGU75B02F839X"), None)
ok("persona fisica in anagrafica senza P.IVA", pf is not None and pf["piva"] == "" and pf["ragione_sociale"] == "VERDI LUIGI", str(pf))

# --- fattura a mano con nuovo cliente e PDF caricato ---
pdf_mano = fattura_finta("_mano.pdf", "FPR 300/26", "05/09/2026", "200,00")
contenuto = pdf_mano.read_bytes()
pdf_mano.unlink()
controlla("fattura a mano: pagina", cl.get("/fatture/nuova"), contiene="NUOVO CLIENTE")
controlla("fattura a mano con nuovo cliente", cl.post("/fatture/nuova", data={
    "codice": "nuovo", "nuovo_ragione_sociale": "bianchi", "nuovo_nome": "anna",
    "nuovo_piva": "", "nuovo_cf": "BNCNNA90C41F839K", "numero": "FPR 300/26",
    "data": "05/09/2026", "scadenza": "05/10/2026", "totale": "200,00",
    "pdf": (__import__("io").BytesIO(contenuto), "scansione.pdf")},
    content_type="multipart/form-data", follow_redirects=True), contiene="Nuovo cliente creato con codice")
bianchi = next((c for c in motore.leggi_anagrafica(motore.apri_excel()) if c["cf"] == "BNCNNA90C41F839K"), None)
ok("cliente a mano creato (BIANCHI ANNA)", bianchi is not None and bianchi["ragione_sociale"] == "BIANCHI"
   and bianchi["nome"] == "ANNA", str(bianchi))
pdf_bianchi = motore.cartella_di(bianchi) / "FT FPR 300-26 del 05-09-2026.pdf" if bianchi else None
ok("PDF caricato salvato col nome standard nella cartella del cliente", pdf_bianchi is not None and pdf_bianchi.exists())
controlla("fattura a mano doppia rifiutata", cl.post("/fatture/nuova", data={
    "codice": str(bianchi["codice"]), "numero": "FPR 300/26", "data": "05/09/2026",
    "scadenza": "", "totale": "200,00"}, content_type="multipart/form-data", follow_redirects=True),
    contiene="Già registrata")
controlla("nuovo cliente con CF già esistente rifiutato", cl.post("/fatture/nuova", data={
    "codice": "nuovo", "nuovo_ragione_sociale": "altro", "nuovo_cf": "BNCNNA90C41F839K",
    "numero": "FPR 301/26", "data": "06/09/2026", "totale": "10,00"},
    content_type="multipart/form-data", follow_redirects=True), contiene="Esiste già il cliente")
controlla("fattura a mano senza PDF su cliente esistente", cl.post("/fatture/nuova", data={
    "codice": str(bianchi["codice"]), "numero": "FPR 302/26", "data": "07/09/2026",
    "scadenza": "07/10/2026", "totale": "30,00"}, content_type="multipart/form-data",
    follow_redirects=True), contiene="Nessun PDF allegato")
controlla("file non PDF rifiutato", cl.post("/fatture/nuova", data={
    "codice": str(bianchi["codice"]), "numero": "FPR 303/26", "data": "07/09/2026", "totale": "30,00",
    "pdf": (__import__("io").BytesIO(b"ciao"), "x.pdf")}, content_type="multipart/form-data",
    follow_redirects=True), contiene="non è un PDF")
controlla("pagamento sulla fattura caricata a mano", cl.post("/pagamenti/registra", data={
    "codice": str(bianchi["codice"]), "importo": "200,00", "data": "", "fattura": "FPR 300/26",
    "modalita": "Contanti"}, follow_redirects=True), contiene="SALDATA")
ok("PDF caricato a mano timbrato e archiviato in SALDATE",
   (motore.SALDATE / f"BIANCHI ANNA - {pdf_bianchi.name}").exists()
   and "PAGATO" in testo_pdf(motore.SALDATE / f"BIANCHI ANNA - {pdf_bianchi.name}"))
fatt_bianchi = [f for f in motore.leggi_fatture(motore.apri_excel()) if f["codice"] == bianchi["codice"]]
ok("BIANCHI: 2 fatture, una saldata e una aperta da 30",
   len(fatt_bianchi) == 2 and sorted(f["residuo"] for f in fatt_bianchi) == [0.0, 30.0])

# --- nuovo cliente, report ---
controlla("nuovo cliente", cl.post("/clienti/salva", data={
    "codice": "", "ragione_sociale": "cliente di prova srl", "nome": "",
    "piva": "12345678901", "cf": "12345678901", "note": "test"},
    follow_redirects=True), contiene="salvato")
r = cl.get("/report")
ok(f"report PDF ({len(r.data)} byte)", r.status_code == 200 and r.data[:4] == b"%PDF")
with pdfplumber.open(__import__("io").BytesIO(r.data)) as f:
    testo_report = "\n".join(p.extract_text() or "" for p in f.pages)
ok("report: tabella unica alfabetica con totale generale", "ordine alfabetico" in testo_report
   and "TOTALE GENERALE" in testo_report and "Situazione per cliente" not in testo_report
   and testo_report.index("ALFA ESEMPIO S.R.L") < testo_report.index("BIANCHI"), testo_report[:200])
ok("report: subtotale del cliente con più fatture", "Totale ALFA ESEMPIO S.R.L" in testo_report and "1.915,00" in testo_report)

# ======== ordine alfabetico e guardia nomi (22/09/2026) ========
corpo = controlla("fatture in ordine alfabetico per cliente", cl.get("/fatture"), contiene="ordine alfabetico")
ok("ALFA prima di BIANCHI nella tabella fatture", corpo.index("20 - ALFA ESEMPIO S.R.L") < corpo.index("- BIANCHI"))
corpo = cl.get("/clienti").get_data(as_text=True)
ok("anagrafica in ordine alfabetico (non per codice)", corpo.index("BIANCHI") < corpo.index("CLIENTE DI PROVA SRL"))
corpo = cl.get("/").get_data(as_text=True)
ok("riepilogo: aperte in ordine alfabetico", corpo.index("ALFA ESEMPIO S.R.L") < corpo.index("BIANCHI"))
controlla("nuovo cliente con nome troncato di uno esistente rifiutato", cl.post("/clienti/salva", data={
    "codice": "", "ragione_sociale": "bianc", "nome": "", "piva": "", "cf": "", "note": ""},
    follow_redirects=True), contiene="Nome troppo simile")
controlla("fattura a mano con nuovo cliente dal nome troncato rifiutata", cl.post("/fatture/nuova", data={
    "codice": "nuovo", "nuovo_ragione_sociale": "BIANCHI AN", "numero": "FPR 777/26", "data": "10/09/2026",
    "totale": "10,00"}, content_type="multipart/form-data", follow_redirects=True), contiene="Nome troppo simile")
controlla("stesso nome ma con P.IVA diversa accettato", cl.post("/clienti/salva", data={
    "codice": "", "ragione_sociale": "bianchi", "nome": "anna", "piva": "33333333333", "cf": "", "note": ""},
    follow_redirects=True), contiene="salvato")

# --- analisi cartelle: PDF tolto dalla cartella cliente -> fattura (e pagamento) eliminati ---
fattura_finta("_sync_a.pdf", "FPR 500/26", "10/09/2026", "600,00")
fattura_finta("_sync_b.pdf", "FPR 501/26", "11/09/2026", "400,00")
controlla("registro 2 fatture per l'analisi", cl.post("/smistamento/conferma", follow_redirects=True),
          contiene="600,00")
controlla("acconto 100 su FPR 500/26", cl.post("/pagamenti/registra", data={
    "codice": CLIENTE, "importo": "100", "data": "", "fattura": "FPR 500/26", "modalita": "Bonifico"},
    follow_redirects=True), contiene="acconto")
cartella_alfa = motore.cartella_di({"codice": int(CLIENTE), "ragione_sociale": "ALFA ESEMPIO S.R.L", "nome": None})
pdf_500 = cartella_alfa / "FT FPR 500-26 del 10-09-2026.pdf"
ok("PDF FPR 500/26 nella cartella (timbrato)", pdf_500.exists() and "ACCONTO" in testo_pdf(pdf_500))
orig_500 = motore.ORIGINALI / cartella_alfa.name / pdf_500.name
ok("originale di FPR 500/26 salvato", orig_500.exists())
n_acconti_prima = len(motore.leggi_acconti(motore.apri_excel()))
corpo = controlla("analisi con tutto allineato", cl.post("/smistamento/analizza"), contiene="Tutto allineato")
ok("nessuna fattura da eliminare se i PDF ci sono", "Fatture da eliminare" not in corpo)
# tolgo il PDF a mano e metto un PDF orfano nella cartella
pdf_500.unlink()
orfano = cartella_alfa / "orfano.pdf"
orfano.write_bytes(b"%PDF-1.4 finto")
corpo = controlla("analisi: fattura con PDF tolto", cl.post("/smistamento/analizza"), contiene="Fatture da eliminare")
ok("anteprima elenca FPR 500/26 con il pagamento da annullare",
   "FPR 500/26" in corpo and "annullato" in corpo and "100,00" in corpo)
ok("anteprima segnala il PDF orfano", "orfano.pdf" in corpo and "non registrati" in corpo)
ok("FPR 501/26 (PDF presente) non è tra quelle da eliminare",
   corpo.count("FPR 501/26") == 0)
# la stessa fattura rimessa nella INBOX non è un doppione: la vecchia va via, la nuova entra
fattura_finta("_sync_a_bis.pdf", "FPR 500/26", "10/09/2026", "600,00")
corpo = controlla("analisi: fattura tolta e rimessa nella INBOX", cl.post("/smistamento/analizza"), contiene="pronta")
ok("la fattura rimessa nella INBOX non è segnata come già registrata", "Già registrata" not in corpo)
corpo = controlla("conferma: elimina + registra", cl.post("/smistamento/conferma", follow_redirects=True),
                  contiene="ELIMINATA FT FPR 500/26")
ok("messaggio di pagamento annullato", "Pagamento di € 100,00" in corpo and "annullato" in corpo)
f500 = [f for f in fatture_cliente() if f["numero"] == "FPR 500/26"]
ok("FPR 500/26 presente una sola volta, riregistrata pulita (residuo 600)",
   len(f500) == 1 and f500[0]["residuo"] == 600.0 and "acconto" not in f500[0]["nota"], str(f500))
ok("pagamento collegato eliminato dal registro Acconti",
   len(motore.leggi_acconti(motore.apri_excel())) == n_acconti_prima - 1)
ok("originale della fattura eliminata rimosso", not orig_500.exists())
ok("PDF nuovo al posto di quello tolto", pdf_500.exists() and "ACCONTO" not in testo_pdf(pdf_500))
f501 = [f for f in fatture_cliente() if f["numero"] == "FPR 501/26"]
ok("FPR 501/26 intatta (residuo 400)", len(f501) == 1 and f501[0]["residuo"] == 400.0)
orfano.unlink()
controlla("analisi di nuovo allineata", cl.post("/smistamento/analizza"), contiene="Tutto allineato")
controlla("conferma senza nulla da fare", cl.post("/smistamento/conferma", follow_redirects=True),
          contiene="Niente da fare")

# ======== correzioni del 22/09/2026 ========
# 1) escape HTML: nome/note con < > & non devono finire grezzi nella pagina
controlla("cliente con caratteri speciali", cl.post("/clienti/salva", data={
    "codice": "", "ragione_sociale": "prova <b>x</b> & c", "nome": "", "piva": "22222222222",
    "cf": "", "note": "<i>n</i>"}, follow_redirects=True), contiene="salvato")
corpo = cl.get("/clienti").get_data(as_text=True)
ok("markup nel nome cliente neutralizzato (escape HTML)",
   "&lt;B&gt;X&lt;/B&gt; &amp; C" in corpo and "<B>X</B>" not in corpo)
ok("il simbolo euro resta leggibile con l'escape attivo", "&euro; " in corpo and "&amp;euro;" not in corpo)

# 2) rinumerazione di una fattura con pagamento: il pagamento segue il nuovo numero
fattura_finta("_rinum.pdf", "FPR 900/26", "01/09/2026", "500,00")
cl.post("/smistamento/conferma", follow_redirects=True)
controlla("acconto 100 su FPR 900/26", cl.post("/pagamenti/registra", data={
    "codice": CLIENTE, "importo": "100", "data": "", "fattura": "FPR 900/26", "modalita": ""},
    follow_redirects=True), contiene="acconto")
f900 = next(f for f in fatture_cliente() if f["numero"] == "FPR 900/26")
controlla("rinumero FPR 900/26 -> FPR 901/26", cl.post(f"/fatture/modifica/{f900['riga']}", data={
    "numero": "FPR 901/26", "data": "01/09/2026", "scadenza": "", "totale": "500,00"},
    follow_redirects=True), contiene="1 pagamento collegato")
acc = next(a for a in motore.leggi_acconti(motore.apri_excel()) if "FPR 90" in a["nota"])
ok("nota del pagamento aggiornata col nuovo numero", "FT FPR 901/26 acconto € 100,00" in acc["nota"], acc["nota"])

# 3) eliminazione di una fattura con pagamenti collegati: rifiutata (niente pagamenti orfani)
f901 = next(f for f in fatture_cliente() if f["numero"] == "FPR 901/26")
controlla("elimina fattura con pagamento: rifiutata", cl.post(f"/fatture/elimina/{f901['riga']}",
          follow_redirects=True), contiene="pagamenti collegati")
ok("fattura ancora presente dopo il rifiuto", any(f["numero"] == "FPR 901/26" for f in fatture_cliente()))
corpo = controlla("annullo del pagamento rinumerato", cl.post(f"/pagamenti/elimina/{acc['riga']}",
                  follow_redirects=True), contiene="timbri tolti")
ok("nessun avviso 'riga non trovata' all'annullo", "riga non trovata" not in corpo)
f901 = next(f for f in fatture_cliente() if f["numero"] == "FPR 901/26")
ok("residuo ripristinato a 500 dopo la rinumerazione", f901["residuo"] == 500.0, str(f901["residuo"]))

# 4) righe scritte a mano male nell'Excel (data come testo, codice non numerico) non bloccano l'app
wb = motore.apri_excel()
ws = wb["Insoluti"]
r_an = motore.prima_riga_libera(ws)
ws.cell(row=r_an, column=1).value = int(CLIENTE)
ws.cell(row=r_an, column=4).value = "31/12/2026"
ws.cell(row=r_an, column=5).value = 10
ws.cell(row=r_an, column=8).value = "FT ANOM del 01/01/2026 — Totale € 10,00"
ws.cell(row=r_an + 1, column=1).value = "abc"
ws.cell(row=r_an + 1, column=5).value = 5
wb.save(motore.EXCEL)
controlla("dashboard con data-testo e codice non numerico nell'Excel", cl.get("/"), contiene="FT ANOM")
f_an = next(f for f in fatture_cliente() if f["numero"] == "ANOM")
ok("data scritta come testo letta come scadenza", f_an["scadenza"] == datetime(2026, 12, 31), str(f_an["scadenza"]))
ok("riga con codice non numerico ignorata", not any(f["riga"] == r_an + 1 for f in motore.leggi_fatture(motore.apri_excel())))
wb = motore.apri_excel()
ws = wb["Insoluti"]
for r in (r_an, r_an + 1):
    for col in (1, 3, 4, 5, 8):
        ws.cell(row=r, column=col).value = None
wb.save(motore.EXCEL)

# 5) registri ampliati: le formule del foglio Clienti seguono il nuovo limite
wb = motore.apri_excel()
formula_j = wb["Clienti"].cell(row=5, column=10).value or ""
ok("formule Clienti riallineate al nuovo limite dei registri", f"$B${motore.ULTIMA_RIGA_REGISTRI}" in formula_j, formula_j)

# 6) cartella cliente sparita (rinominata per sbaglio): segnalata, nessuna eliminazione proposta
cartella_rinominata = cartella_alfa.with_name(cartella_alfa.name + " - rinominata")
cartella_alfa.rename(cartella_rinominata)
corpo = controlla("analisi con cartella cliente sparita", cl.post("/smistamento/analizza"),
                  contiene="Cartelle cliente non trovate")
ok("nessuna fattura da eliminare per una cartella sparita", "Fatture da eliminare" not in corpo)
ok("fatture di ALFA ancora registrate", len(fatture_cliente()) >= 1)
cartella_rinominata.rename(cartella_alfa)
controlla("analisi allineata dopo il ripristino della cartella", cl.post("/smistamento/analizza"),
          contiene="Tutto allineato")

# 7) PDF bloccato da un altro programma durante lo spostamento: niente errore 500, messaggio chiaro
p_bloccato = fattura_finta("_bloccato.pdf", "FPR 910/26", "03/09/2026", "70,00")
handle = open(p_bloccato, "rb")   # su Windows impedisce lo spostamento del file
corpo = controlla("conferma con PDF bloccato", cl.post("/smistamento/conferma", follow_redirects=True),
                  contiene="NON spostato")
handle.close()
ok("fattura registrata comunque", any(f["numero"] == "FPR 910/26" for f in fatture_cliente()))
ok("PDF rimasto nella INBOX", p_bloccato.exists())
controlla("analisi dopo il blocco propone la sistemazione", cl.post("/smistamento/analizza"),
          contiene="Fatture da eliminare")
controlla("conferma: fattura riregistrata pulita", cl.post("/smistamento/conferma", follow_redirects=True),
          contiene="ELIMINATA FT FPR 910/26")
ok("FPR 910/26 una sola volta e PDF nella cartella",
   len([f for f in fatture_cliente() if f["numero"] == "FPR 910/26"]) == 1
   and (cartella_alfa / "FT FPR 910-26 del 03-09-2026.pdf").exists())

# ======== aspetto (versione bianca: nessun logo) ========
for r in ("/risorse/Montserrat.woff2",):
    ok(f"risorsa servita dall'app: {r}", cl.get(r).status_code == 200)
corpo = controlla("riepilogo con grafico mensile", cl.get("/"), contiene="Andamento ultimi 12 mesi")
ok("grafico SVG con barre, legenda e tooltip", 'class="grafico"' in corpo and 'class="mese"' in corpo
   and "Incassato" in corpo and "data-tip=" in corpo)
ok("titolo scheda generico e favicon", "<title>Riepilogo – Scadenziario</title>" in corpo and 'rel="icon"' in corpo)
ok("nessun logo, nome studio in barra e menu con icone", "<img" not in corpo and "logo_" not in corpo and "Studio di prova Esempio" in corpo
   and corpo.count("<svg") >= 9 and 'href="/impostazioni"' in corpo)
ok("font Montserrat incluso", "Montserrat.woff2" in corpo)
ok("nessun campo password e nessun pulsante Esci", 'type="password"' not in corpo and "/esci" not in corpo)
corpo = controlla("fatture con ricerca e filtri", cl.get("/fatture"), contiene='data-filtra="tab-fatture"')
ok("filtri stato e anno sulle fatture", 'data-campo="stato"' in corpo and 'data-campo="anno"' in corpo
   and 'data-stato="Aperta"' in corpo)
controlla("pagamenti con ricerca", cl.get("/pagamenti"), contiene='data-filtra="tab-pagamenti"')
corpo = controlla("clienti con ricerca", cl.get("/clienti"), contiene='data-filtra="tab-clienti"')
ok("clienti: filtro per situazione", 'data-situazione="aperto"' in corpo and 'data-campo="situazione"' in corpo)
and_ = motore.andamento_mensile([{"data": "05/09/2026", "totale": 100.0, "residuo": 40.0},
                                 {"data": "05/09/2025", "totale": 999.0, "residuo": 0.0}], oggi=datetime(2026, 9, 22))
ok("andamento mensile: 12 mesi, ultimo = set 26 con 100/60/40, fuori finestra ignorato",
   len(and_) == 12 and and_[-1]["etichetta"] == "set 26" and and_[-1]["fatturato"] == 100.0
   and and_[-1]["incassato"] == 60.0 and and_[-1]["da_incassare"] == 40.0 and and_[0]["fatturato"] == 0.0, str(and_[-1]))
graf = motore.grafico_mensile(and_)
ok("grafico: barra impilata con spazio di 2px e tacche pulite",
   graf["barre"][-1]["h_inc"] > 0 and graf["barre"][-1]["h_da"] > 0
   and graf["barre"][-1]["y_da"] + graf["barre"][-1]["h_da"] + 2 == graf["barre"][-1]["y_inc"]
   and graf["griglia"][-1]["val"] == 200, str(graf["barre"][-1]))
r = cl.get("/report")
with pdfplumber.open(__import__("io").BytesIO(r.data)) as f:
    testo_report = "\n".join(p.extract_text() or "" for p in f.pages)
    n_img = len(f.pages[0].images)
ok("report PDF con intestazione dello studio (dalle Impostazioni), senza logo", "Studio di prova Esempio" in testo_report
   and "P.IVA 01234567897" in testo_report and n_img == 0, testo_report[:120])

# ======== scheda cliente (23/09/2026) ========
corpo = controlla("pagina Clienti con tendina 'Apri scheda'", cl.get("/clienti"), contiene='id="scegli-scheda"')
ok("nome cliente cliccabile e pulsante Scheda", f'href="/clienti/{CLIENTE}"' in corpo and ">Scheda<" in corpo)
corpo = controlla("scheda del cliente ALFA", cl.get(f"/clienti/{CLIENTE}"), contiene="ALFA ESEMPIO S.R.L")
ok("scheda: riquadri, fatture e pagamenti", "Fatture registrate" in corpo and "Da incassare" in corpo
   and "Pagamenti ricevuti" in corpo and "Scarica scheda PDF" in corpo and 'href="/clienti/20/pdf"' in corpo)
sc = motore.scheda_cliente(int(CLIENTE))
ok("scheda_cliente: fatture e totali coerenti con l'Excel",
   sc["n_fatture"] == len(fatture_cliente()) and abs(sc["fatturato"] - sc["incassato"] - sc["da_incassare"] + sc["credito"]) < 0.01
   and sc["n_aperte"] == len([f for f in fatture_cliente() if f["residuo"] > 0]), str({k: sc[k] for k in ("n_fatture", "fatturato", "incassato", "da_incassare")}))
ok("scheda: tutte le fatture della pagina sono del cliente", all(f"FT {f['numero']}" in corpo for f in fatture_cliente()))
corpo = controlla("scheda cliente senza fatture", cl.get(f"/clienti/{SENZA_FATTURE}"), contiene="Nessuna fattura registrata")
corpo = controlla("scheda cliente inesistente -> torna a Clienti", cl.get("/clienti/999999", follow_redirects=True),
                  contiene="non trovato")
r = cl.get(f"/clienti/{CLIENTE}/pdf")
ok("PDF scheda scaricato", r.status_code == 200 and r.data[:4] == b"%PDF" and "attachment" in r.headers.get("Content-Disposition", ""))
with pdfplumber.open(__import__("io").BytesIO(r.data)) as f:
    testo_scheda = " ".join(p.extract_text() or "" for p in f.pages)
ok("PDF scheda: intestazione studio, dati cliente, fatture aperte, senza testo finale",
   "Studio di prova Esempio" in testo_scheda and "ALFA ESEMPIO S.R.L" in testo_scheda and "07654321004" in testo_scheda
   and "Fatture aperte" in testo_scheda and "Pagamenti ricevuti" in testo_scheda
   and "Gentile cliente" not in testo_scheda, testo_scheda[:300])
pdf_schede = list(motore.SCHEDE.glob("20 - *scheda*.pdf"))
ok("PDF scheda salvato in dati/SCHEDE CLIENTI", len(pdf_schede) == 1, str(pdf_schede))
r = cl.get("/report")
ok("report generale ancora funzionante dopo il refactoring dell'intestazione", r.status_code == 200 and r.data[:4] == b"%PDF")

# ======== versione bianca: Impostazioni e primo avvio ========
corpo = controlla("pagina Impostazioni", cl.get("/impostazioni"), contiene="Dati dello studio")
ok("impostazioni mostrano i dati salvati", 'value="01234567897"' in corpo and "Studio di prova Esempio" in corpo)
controlla("salvo impostazioni", cl.post("/impostazioni", data={
    "nome": "Studio di prova Esempio", "piva": "1234567897", "cf": "prvstd70a01h501z", "indirizzo": "Via Prova 1 - Roma",
    "telefono": "000 000", "email": "prova@esempio.it", "intestazione_fatture": ""},
    follow_redirects=True), contiene="Impostazioni salvate")
cfg = motore.carica_config()
ok("P.IVA normalizzata, CF maiuscolo, nessuna password in config", cfg["piva_studio"] == "01234567897"
   and cfg["cf_studio"] == "PRVSTD70A01H501Z" and "password" not in cfg, str(cfg))
ok("impostazioni senza campo password", 'type="password"' not in cl.get("/impostazioni").get_data(as_text=True))
import tempfile
with tempfile.TemporaryDirectory() as tmp:
    nuova = Path(tmp) / "dati"
    motore.inizializza_dati(nuova)
    ok("primo avvio: cartella dati vuota creata", motore.raggiungibile(nuova)
       and all((nuova / c).is_dir() for c in motore.CARTELLE_DATI) and (nuova / "config.json").exists())
    from openpyxl import load_workbook
    wb_vuoto = load_workbook(nuova / motore.NOME_EXCEL)
    salvato, motore.EXCEL = motore.EXCEL, nuova / motore.NOME_EXCEL
    ok("modello vuoto: zero clienti, zero fatture, zero pagamenti", motore.leggi_anagrafica(wb_vuoto) == []
       and motore.leggi_fatture(wb_vuoto) == [] and motore.leggi_acconti(wb_vuoto) == [])
    motore.EXCEL = salvato
    ok("rete non raggiungibile: nessuna cartella creata", motore.inizializza_dati(Path(tmp) / "manca" / "dati") is None
       and not (Path(tmp) / "manca").exists())

# --- pulizia: elimino le fatture di prova ---
for fatt in fatture_cliente():
    controlla(f"elimina fattura riga {fatt['riga']}", cl.post(f"/fatture/elimina/{fatt['riga']}",
              follow_redirects=True), contiene="eliminata")
cli = [c for c in motore.leggi_anagrafica(motore.apri_excel()) if c["codice"] == int(CLIENTE)][0]
ok("dovuto ALFA azzerato", (cli["dovuto"] or 0) == 0, str(cli["dovuto"]))
n_backup = len(list(motore.BACKUP.glob("backup_*.xlsx")))
ok(f"backup automatici creati: {n_backup}", n_backup >= 10)

print(f"\nRISULTATO: {sum(esiti)} controlli superati su {len(esiti)}"
      + ("" if all(esiti) else "  <-- CI SONO FALLIMENTI"))
sys.exit(0 if all(esiti) else 1)
